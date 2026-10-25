"""Informe en Excel para revisar el periodo antes de aprobar las facturas."""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .comisiones import Calculo
from .modelo import Incidencia, Movimiento, Posicion

EUR = '#,##0.00 "€"'
CABECERA = PatternFill("solid", fgColor="1D3557")
ALTA = PatternFill("solid", fgColor="FDE2E2")
MEDIA = PatternFill("solid", fgColor="FEF3C7")
LISTA = PatternFill("solid", fgColor="DCFCE7")


def _hoja(libro, titulo, cabeceras, filas, anchos, columnas_eur=(), relleno=None):
    hoja = libro.create_sheet(titulo)
    hoja.append(cabeceras)
    for celda in hoja[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = CABECERA
        celda.alignment = Alignment(vertical="center", wrap_text=True)
    for fila in filas:
        hoja.append(fila)
        if relleno and (color := relleno(fila)):
            for celda in hoja[hoja.max_row]:
                celda.fill = color
    for i, ancho in enumerate(anchos, start=1):
        hoja.column_dimensions[get_column_letter(i)].width = ancho
    for col in columnas_eur:
        for celda in hoja[get_column_letter(col)][1:]:
            celda.number_format = EUR
    for fila in hoja.iter_rows(min_row=2):
        for celda in fila:
            celda.alignment = Alignment(vertical="top", wrap_text=True)
    hoja.freeze_panes = "A2"
    hoja.auto_filter.ref = hoja.dimensions
    return hoja


def generar(ruta: Path, periodo_nombre: str, calculos: list[Calculo], estados: dict[str, str],
            incidencias: list[Incidencia], posiciones: list[Posicion], movimientos: list[Movimiento],
            numeros: dict[str, str]):
    libro = Workbook()
    libro.remove(libro.active)

    por_cliente = {c.cliente: [i for i in incidencias if i.cliente == c.cliente] for c in calculos}
    _hoja(libro, "Resumen",
          ["Cliente", "Nombre", "Patrimonio validado", "Base imponible", "IVA", "Total factura",
           "Incidencias altas", "Incidencias medias", "Estado", "Factura (borrador)"],
          [[c.cliente, c.nombre, c.patrimonio, c.base_imponible, c.cuota_iva, c.total,
            sum(i.severidad == "alta" for i in por_cliente[c.cliente]),
            sum(i.severidad == "media" for i in por_cliente[c.cliente]),
            estados[c.cliente], numeros[c.cliente]] for c in calculos]
          + [["TOTAL", "", sum(c.patrimonio for c in calculos), sum(c.base_imponible for c in calculos),
              sum(c.cuota_iva for c in calculos), sum(c.total for c in calculos), "", "", "", ""]],
          [9, 30, 18, 16, 13, 16, 11, 11, 20, 20], columnas_eur=(3, 4, 5, 6),
          relleno=lambda f: LISTA if f[8] == "Lista para aprobar" else (MEDIA if f[8] else None))
    libro["Resumen"].insert_rows(1)
    libro["Resumen"]["A1"] = f"Conciliación y comisiones · {periodo_nombre} · datos ficticios"
    libro["Resumen"]["A1"].font = Font(bold=True, size=13, color="1D3557")
    libro["Resumen"].freeze_panes = "A3"
    libro["Resumen"].auto_filter.ref = None

    _hoja(libro, "Incidencias",
          ["Severidad", "Cliente", "Cuenta", "Tipo", "ISIN", "Descripción", "Banco", "Interno",
           "Diferencia", "Qué hacer"],
          [[i.severidad.upper(), i.cliente, i.cuenta, i.tipo, i.isin, i.descripcion,
            i.valor_banco, i.valor_interno, i.diferencia, i.accion]
           for i in sorted(incidencias, key=lambda i: (i.severidad != "alta", i.cliente))],
          [10, 9, 16, 26, 15, 60, 14, 14, 14, 42], columnas_eur=(7, 8, 9),
          relleno=lambda f: ALTA if f[0] == "ALTA" else MEDIA)

    _hoja(libro, "Cálculo comisiones", ["Cliente", "Paso del cálculo"],
          [[c.cliente, linea] for c in calculos for linea in c.desglose], [9, 120])

    _hoja(libro, "Patrimonio validado",
          ["Cliente", "Cuenta", "ISIN", "Nombre", "Títulos", "Precio", "Divisa", "Tipo cambio",
           "Valor EUR", "Fecha", "Fuente"],
          [[p.cliente, p.cuenta, p.isin, p.nombre, p.titulos, p.precio, p.divisa, p.tipo_cambio,
            p.valor_eur, p.fecha, p.fuente] for p in sorted(posiciones, key=lambda p: (p.cliente, p.isin))],
          [9, 16, 15, 34, 12, 12, 8, 11, 15, 12, 40], columnas_eur=(9,))

    _hoja(libro, "Movimientos banco",
          ["Cliente", "Cuenta", "Fecha", "Tipo", "ISIN", "Importe EUR", "Concepto", "Fuente"],
          [[m.cliente, m.cuenta, m.fecha, m.tipo, m.isin, m.importe_eur, m.descripcion, m.fuente]
           for m in sorted(movimientos, key=lambda m: (m.cliente, m.fecha))],
          [9, 16, 12, 13, 15, 15, 28, 40], columnas_eur=(6,))

    for hoja in libro.worksheets:
        for fila in hoja.iter_rows():
            for celda in fila:
                if hasattr(celda.value, "year"):
                    celda.number_format = "DD/MM/YYYY"
    libro.save(ruta)
