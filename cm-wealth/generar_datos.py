"""Genera los datos ficticios de ejemplo en datos/.

El registro interno lleva errores a propósito, para que la conciliación tenga algo que encontrar:
  C001  precio desactualizado del Euro Bond Fund; compra del ETF Mundo y un dividendo sin registrar
  C002  venta parcial de Banca Levante sin registrar (títulos y efectivo no cuadran)
  C003  tipo de cambio USD y efectivo ligeramente distintos, dentro de la tolerancia (no es incidencia)
"""
import csv
from datetime import date, datetime
from decimal import Decimal as D
from pathlib import Path

from openpyxl import Workbook
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from cm.modelo import euros, formato_num

DATOS = Path(__file__).parent / "datos"
FECHA = date(2026, 9, 30)

# Lo que dice el banco: (cuenta, isin, nombre, títulos, precio, divisa, tipo de cambio)
POSICIONES = {
    "ATL-10045821": [
        ("LU1111111111", "Global Equity Fund A", D("2950.3120"), D("182.4500"), "EUR", D(1)),
        ("IE2222222222", "ETF Mundo Acc", D("240"), D("105.1000"), "EUR", D(1)),
        ("ES0333333333", "Euro Bond Fund", D("3100"), D("112.3600"), "EUR", D(1)),
        ("ES0444444444", "Acciones Energía Ibérica", D("8000"), D("14.2350"), "EUR", D(1)),
    ],
    "ATL-10046003": [
        ("LU6666666666", "Fondo Mixto Flexible", D("6800"), D("98.4000"), "EUR", D(1)),
        ("ES0777777777", "Acciones Banca Levante", D("21500"), D("7.8840"), "EUR", D(1)),
        ("FR8888888888", "Fondo Monetario Euro", D("3050"), D("141.2200"), "EUR", D(1)),
    ],
    "MER-77310459": [
        ("US9999999999", "Acciones Tech Norteamérica", D("100"), D("230.5000"), "USD", D("1.0850")),
        ("IE1212121212", "ETF Bonos Gobierno Euro", D("900"), D("52.1800"), "EUR", D(1)),
        ("LU3434343434", "Fondo Small Caps Europa", D("410"), D("121.7000"), "EUR", D(1)),
    ],
}
EFECTIVO = {"ATL-10045821": D("58420.37"), "ATL-10046003": D("21038.40"), "MER-77310459": D("6712.45")}
TITULARES = {"ATL-10045821": ("C001", "Laura Martín Ruiz"), "ATL-10046003": ("C002", "Inversiones Alameda SL"),
             "MER-77310459": ("C003", "Jorge Ortega Vidal")}

# Movimientos del trimestre según el banco: (cuenta, fecha, concepto, isin, importe)
MOVIMIENTOS = [
    ("ATL-10045821", date(2026, 7, 1), "COMISION CUSTODIA", "", D("-120.50")),
    ("ATL-10045821", date(2026, 7, 28), "DIVIDENDO NETO", "ES0444444444", D("1944.00")),
    ("ATL-10045821", date(2026, 9, 12), "COMPRA", "IE2222222222", D("-25200.00")),
    ("ATL-10046003", date(2026, 7, 1), "COMISION CUSTODIA", "", D("-185.20")),
    ("ATL-10046003", date(2026, 8, 3), "APORTACION", "", D("20000.00")),
    ("ATL-10046003", date(2026, 9, 22), "VENTA", "ES0777777777", D("11850.00")),
    ("MER-77310459", date(2026, 8, 17), "APORTACION", "", D("125000.00")),
    ("MER-77310459", date(2026, 8, 18), "COMPRA", "IE1212121212", D("-46800.00")),
    ("MER-77310459", date(2026, 8, 18), "COMPRA", "LU3434343434", D("-49610.00")),
    ("MER-77310459", date(2026, 8, 19), "COMPRA", "US9999999999", D("-21877.55")),
]


def valor(titulos, precio, cambio):
    return euros(titulos * precio / cambio)


def es(n, decimales=2):
    return formato_num(n, decimales)


def pdf_banco_atlantico(ruta: Path):
    estilos = getSampleStyleSheet()
    doc = SimpleDocTemplate(str(ruta), pagesize=A4, leftMargin=15 * mm, rightMargin=15 * mm,
                            title="Extracto de posiciones", author="Banco Atlántico (ficticio)")
    estilo = TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                         ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0b3d5c")),
                         ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("FONTSIZE", (0, 0), (-1, -1), 8.5),
                         ("ALIGN", (2, 1), (-1, -1), "RIGHT")])
    e = []
    for cuenta in ("ATL-10045821", "ATL-10046003"):
        _, titular = TITULARES[cuenta]
        e += [Paragraph("<b>BANCO ATLÁNTICO</b> · Extracto de posiciones (ficticio)", estilos["Title"]),
              Paragraph(f"Cuenta de valores: {cuenta}", estilos["Normal"]),
              Paragraph(f"Titular: {titular}", estilos["Normal"]),
              Paragraph(f"Fecha de valoración: {FECHA:%d/%m/%Y}", estilos["Normal"]), Spacer(1, 6 * mm)]
        filas = [["ISIN", "Descripción", "Títulos", "Precio", "Divisa", "Valor (EUR)"]]
        total = D(0)
        for isin, nombre, tit, precio, divisa, cambio in POSICIONES[cuenta]:
            v = valor(tit, precio, cambio)
            total += v
            filas.append([isin, nombre, es(tit, 4 if tit % 1 else 0), es(precio, 4), divisa, es(v)])
        filas.append(["—", "Efectivo (saldo cuenta)", "", "", "EUR", es(EFECTIVO[cuenta])])
        total += EFECTIVO[cuenta]
        filas.append(["TOTAL", "", "", "", "", es(total)])
        t = Table(filas, colWidths=[28 * mm, 58 * mm, 22 * mm, 24 * mm, 14 * mm, 30 * mm])
        t.setStyle(estilo)
        e += [t, Spacer(1, 8 * mm), Paragraph("Movimientos del periodo 01/07/2026 – 30/09/2026", estilos["Heading3"])]
        movs = [["Fecha", "Concepto", "ISIN", "Importe (EUR)"]]
        movs += [[f"{f:%d/%m/%Y}", c, i or "—", es(imp)] for cu, f, c, i, imp in MOVIMIENTOS if cu == cuenta]
        t = Table(movs, colWidths=[28 * mm, 70 * mm, 40 * mm, 38 * mm])
        t.setStyle(estilo)
        e += [t, PageBreak()]
    doc.build(e[:-1])


def csv_banco_meridional(ruta_pos: Path, ruta_mov: Path):
    cuenta = "MER-77310459"
    with open(ruta_pos, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(["Cuenta", "ISIN", "Descripción", "Cantidad", "Precio", "Divisa", "Cambio EUR",
                    "Valor EUR", "Fecha valoración"])
        for isin, nombre, tit, precio, divisa, cambio in POSICIONES[cuenta]:
            w.writerow([cuenta, isin, nombre, es(tit, 0), es(precio, 4), divisa, es(cambio, 4),
                        es(valor(tit, precio, cambio)), f"{FECHA:%d/%m/%Y}"])
        w.writerow([cuenta, "LIQUIDEZ", "Saldo en euros", "1", es(EFECTIVO[cuenta]), "EUR", "1,0000",
                    es(EFECTIVO[cuenta]), f"{FECHA:%d/%m/%Y}"])
    with open(ruta_mov, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(["Cuenta", "Fecha", "Concepto", "ISIN", "Importe EUR"])
        for cu, fecha, concepto, isin, imp in MOVIMIENTOS:
            if cu == cuenta:
                w.writerow([cu, f"{fecha:%d/%m/%Y}", concepto, isin, es(imp)])


def excel_interno(ruta: Path):
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Posiciones"
    hoja.append(["Cliente", "Cuenta", "ISIN", "Nombre", "Títulos", "Precio", "Divisa", "Tipo cambio",
                 "Valor EUR", "Fecha"])
    efectivo = dict(EFECTIVO)
    for cuenta, posiciones in POSICIONES.items():
        cliente, _ = TITULARES[cuenta]
        for isin, nombre, tit, precio, divisa, cambio in posiciones:
            if isin == "IE2222222222":
                continue  # compra no registrada
            if isin == "ES0333333333":
                precio = D("110.9000")  # precio de hace semanas
            if isin == "ES0777777777":
                tit = D("23000")  # falta la venta de 1.500 títulos
            if isin == "US9999999999":
                cambio = D("1.0870")  # cambio de otro proveedor
            hoja.append([cliente, cuenta, isin, nombre, float(tit), float(precio), divisa, float(cambio),
                         float(valor(tit, precio, cambio)), datetime(2026, 9, 30)])
    efectivo["ATL-10045821"] += D("25200.00") - D("1944.00")
    efectivo["ATL-10046003"] -= D("11850.00")
    efectivo["MER-77310459"] = D("6700.00")
    for cuenta, saldo in efectivo.items():
        hoja.append([TITULARES[cuenta][0], cuenta, "EFECTIVO", "Efectivo", 1, float(saldo), "EUR", 1,
                     float(saldo), datetime(2026, 9, 30)])

    hoja = libro.create_sheet("Movimientos")
    hoja.append(["Cliente", "Cuenta", "Fecha", "Tipo", "ISIN", "Importe EUR", "Descripción"])
    no_registrados = {("ATL-10045821", "ES0444444444"), ("ATL-10045821", "IE2222222222"),
                      ("ATL-10046003", "ES0777777777")}
    tipos = {"COMISION CUSTODIA": "COMISION", "DIVIDENDO NETO": "DIVIDENDO"}
    for cuenta, fecha, concepto, isin, imp in MOVIMIENTOS:
        if (cuenta, isin) in no_registrados:
            continue
        if cuenta == "MER-77310459" and concepto == "APORTACION":
            fecha = date(2026, 8, 15)  # la gestora anota la fecha de la orden; el banco, la de abono
        hoja.append([TITULARES[cuenta][0], cuenta, datetime(fecha.year, fecha.month, fecha.day),
                     tipos.get(concepto, concepto), isin, float(imp), concepto.capitalize()])

    hoja = libro.create_sheet("Historial")
    hoja.append(["Cliente", "Patrimonio inicio periodo", "Fecha", "Marca de agua"])
    hoja.append(["C001", 1052310.18, datetime(2026, 6, 30), None])
    hoja.append(["C002", 1215300.00, datetime(2026, 6, 30), 1238900.00])
    hoja.append(["C003", 0, datetime(2026, 6, 30), None])
    libro.save(ruta)


if __name__ == "__main__":
    DATOS.mkdir(exist_ok=True)
    pdf_banco_atlantico(DATOS / "banco_atlantico_extracto_2026-09-30.pdf")
    csv_banco_meridional(DATOS / "banco_meridional_posiciones_2026-09-30.csv",
                         DATOS / "banco_meridional_movimientos_3T2026.csv")
    excel_interno(DATOS / "cartera_interna.xlsx")
    print(f"Datos ficticios generados en {DATOS}")
